import subprocess
import json
import argparse
import csv
import datetime
import re
from openpyxl import Workbook
from openpyxl.styles import Font

TIER_MAP = {
    "db-f1-micro": (1, 0.614),
    "db-g1-small": (1, 1.7),
    "db-perf-optimized-N-2": (2, 16.0),
    "db-perf-optimized-N-4": (4, 32.0),
    "db-perf-optimized-N-8": (8, 64.0),
    "db-perf-optimized-N-16": (16, 128.0),
    "db-perf-optimized-N-32": (32, 256.0),
    "db-perf-optimized-N-48": (48, 384.0),
    "db-perf-optimized-N-64": (64, 512.0),
    "db-perf-optimized-N-80": (80, 640.0),
    "db-perf-optimized-N-96": (96, 768.0),
    "db-perf-optimized-N-128": (128, 864.0),
    "db-c4a-highmem-2": (2, 16.0),
    "db-c4a-highmem-4": (4, 32.0),
    "db-c4a-highmem-8": (8, 64.0),
    "db-c4a-highmem-16": (16, 128.0),
    "db-c4a-highmem-32": (32, 256.0),
    "db-c4a-highmem-48": (48, 384.0),
    "db-c4a-highmem-64": (64, 512.0),
    "db-c4a-highmem-72": (72, 576.0),
}

def parse_tier_resources(tier):
    """Parses vCPU and memory (in GB) from a Cloud SQL machine tier string."""
    if not tier or tier == "N/A":
        return {"vcpu": "N/A", "memory_gb": "N/A"}

    if tier in TIER_MAP:
        vcpu, memory_gb = TIER_MAP[tier]
        return {"vcpu": vcpu, "memory_gb": memory_gb}

    n4_match = re.match(r"^db-custom-N4-(\d+)-(\d+)$", tier)
    if n4_match:
        vcpu = int(n4_match.group(1))
        memory_mb = int(n4_match.group(2))
        return {"vcpu": vcpu, "memory_gb": round(memory_mb / 1024, 2)}

    custom_match = re.match(r"^db-custom-(\d+)-(\d+)$", tier)
    if custom_match:
        vcpu = int(custom_match.group(1))
        memory_mb = int(custom_match.group(2))
        return {"vcpu": vcpu, "memory_gb": round(memory_mb / 1024, 2)}

    suffix_match = re.search(r"-(\d+)$", tier)
    if suffix_match:
        vcpu = int(suffix_match.group(1))
        return {"vcpu": vcpu, "memory_gb": "Check API"}

    return {"vcpu": "N/A", "memory_gb": "N/A"}

def parse_replica_info(data):
    """Extracts replica count and replica names from gcloud JSON output."""
    replica_list = data.get("replicaNames", [])
    if isinstance(replica_list, list) and replica_list:
        return {
            "replica_count": len(replica_list),
            "replica_names": "\n".join(replica_list)
        }
    return {
        "replica_count": 0,
        "replica_names": "None"
    }

def describe_sql_instance(instance_name, project_id):
    command = [
        "gcloud", "sql", "instances", "describe", instance_name,
        "--project", project_id,
        "--format", "json"
    ]
    try:
        result = subprocess.run(command, capture_output=True, text=True, check=True)
        return json.loads(result.stdout)
    except subprocess.CalledProcessError as e:
        print(f"Error executing command for {instance_name}: {e.stderr}")
        return None

def write_to_excel(final_output, file_path):
    wb = Workbook()
    ws = wb.create_sheet(title="CloudSQL Configs")
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    headers_written = False

    for instance_data in final_output:
        headers = []
        values = []

        for key, value in instance_data.items():
            headers.append(key)
            if isinstance(value, list):
                if key == "IP Address List":
                    clean_value = "\n".join([f"{i['ipAddress']} ({i['type']})" for i in value if 'ipAddress' in i])
                else:
                    clean_value = "\n".join([", ".join([f"{k}: {v}" for k, v in i.items()]) for i in value if isinstance(i, dict)])
                values.append(clean_value)
            elif isinstance(value, dict):
                clean_value = "\n".join([f"{k}: {v}" for k, v in value.items()])
                values.append(clean_value)
            else:
                values.append(value)

        if not headers_written:
            ws.append(headers)
            header_font = Font(bold=True)
            for cell in ws[1]:
                cell.font = header_font
            headers_written = True

        ws.append(values)

    wb.save(file_path)

def helper_function(data, path):
    keys = path.split('.')
    current = data
    for key in keys:
        if isinstance(current, dict):
            current = current.get(key, "N/A")
        else:
            return "N/A"
    return current

def get_nested_value(data, path):
    if isinstance(path, dict):
        return {label: helper_function(data, gcloud_config) for label, gcloud_config in path.items()}
    return helper_function(data, path)

def arrange_data_based_on_json_file(data, template_file):
    with open(template_file, 'r') as f:
        template = json.load(f)

    tier = helper_function(data, "settings.tier")
    resources = parse_tier_resources(tier)
    
    replicas = parse_replica_info(data)

    data['_calculated'] = {
        **resources,
        **replicas
    }

    db_version = str(data.get('databaseInstalledVersion', '')).lower()
    if "mysql" in db_version:
        template['Point-in-Time Recovery'] = "settings.backupConfiguration.binaryLogEnabled"

    final_report = {}
    for label, gcloud_config_path in template.items():
        final_report[label] = get_nested_value(data, gcloud_config_path)

    return final_report

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="CloudSQL Instance Configs Listing script")
    parser.add_argument("--filename_prefix", type=str, required=True, help="Prefix for the output Excel filename")
    parser.add_argument("--config_file", type=str, required=False, default="instance_inventory.csv", help="Path to the inventory CSV")
    parser.add_argument("--mapping_file", type=str, required=False, default="inventory.json", help="Path to the report template JSON")
    args = parser.parse_args()

    final_output = []
    timestamp = datetime.datetime.now().strftime("%Y%m%d%H%M%S")
    filename = f"{args.filename_prefix}_{timestamp}.xlsx"

    with open(args.config_file, newline='') as csvfile:
        reader = csv.DictReader(csvfile)
        for row in reader:
            project = row["project"]
            instance_name = row["instance_name"]
            instance_data = describe_sql_instance(instance_name, project)

            if instance_data:
                instance_arranged_data = arrange_data_based_on_json_file(instance_data, args.mapping_file)
                final_output.append(instance_arranged_data)
                print(f"Successfully retrieved details for: {instance_name}")
            else:
                print(f"Failed to retrieve details for: {instance_name}")

    if final_output:
        write_to_excel(final_output, filename)
        print(f"Report written to: {filename}")
    else:
        print("No instance data retrieved. File creation skipped.")