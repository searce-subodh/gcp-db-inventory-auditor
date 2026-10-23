import json
import argparse
import datetime
import sys
from openpyxl import Workbook
from openpyxl.styles import Font

from googleapiclient import discovery
from googleapiclient.errors import HttpError
import google.auth

def get_active_projects(crm_service):
    """Fetches all active projects accessible by the authenticated credentials."""
    print("Discovering all active projects you have access to...")
    projects = []
    query = "state:ACTIVE"
    
    try:
        request = crm_service.projects().search(query=query)
        while request is not None:
            response = request.execute()
            for project in response.get('projects', []):
                projects.append(project.get('projectId'))
            request = crm_service.projects().search_next(previous_request=request, previous_response=response)
        return projects
    except HttpError as e:
        print(f"Error fetching projects: {e}")
        sys.exit(1)

def calculate_bq_objects_and_owners(bq_service, project_id, dataset_id, dataset_details):
    """Calculates object counts and parses owners to populate the _calculated fields."""
    
    # 1. Parse dataset owners from the 'access' list
    owners = []
    access_list = dataset_details.get("access", [])
    for entry in access_list:
        if entry.get("role") == "OWNER":
            owner = entry.get("userByEmail") or entry.get("specialGroup") or entry.get("groupByEmail")
            if owner:
                owners.append(owner)
    
    dataset_owners = ", ".join(owners) if owners else "None"

    # 2. Get Object Counts using BigQuery API
    table_count = 0
    view_count = 0
    routine_count = 0
    model_count = 0

    try:
        # Tables & Views
        tables_req = bq_service.tables().list(projectId=project_id, datasetId=dataset_id)
        while tables_req is not None:
            tables_resp = tables_req.execute()
            for obj in tables_resp.get('tables', []):
                obj_type = obj.get("type", "")
                if obj_type in ["TABLE", "EXTERNAL", "SNAPSHOT", "CLONE"]:
                    table_count += 1
                elif obj_type in ["VIEW", "MATERIALIZED_VIEW"]:
                    view_count += 1
            tables_req = bq_service.tables().list_next(previous_request=tables_req, previous_response=tables_resp)

        # Routines (Stored Procedures, UDFs)
        routines_req = bq_service.routines().list(projectId=project_id, datasetId=dataset_id)
        while routines_req is not None:
            routines_resp = routines_req.execute()
            routine_count += len(routines_resp.get('routines', []))
            routines_req = bq_service.routines().list_next(previous_request=routines_req, previous_response=routines_resp)

        # Models
        models_req = bq_service.models().list(projectId=project_id, datasetId=dataset_id)
        while models_req is not None:
            models_resp = models_req.execute()
            model_count += len(models_resp.get('models', []))
            models_req = bq_service.models().list_next(previous_request=models_req, previous_response=models_resp)

    except HttpError as e:
        print(f"  -> Warning: Could not list objects for {dataset_id}: {e._get_reason()}")

    total_object_count = table_count + view_count + routine_count + model_count

    return {
        "dataset_owners": dataset_owners,
        "table_count": table_count,
        "view_count": view_count,
        "routine_count": routine_count,
        "model_count": model_count,
        "total_object_count": total_object_count
    }

def get_datasets_for_project(bq_service, project_id):
    """Lists all datasets in a project and fetches their full metadata."""
    datasets_data = []
    try:
        request = bq_service.datasets().list(projectId=project_id)
        while request is not None:
            response = request.execute()
            for ds in response.get('datasets', []):
                dataset_id = ds['datasetReference']['datasetId']
                
                # The list API only returns partial data. Get full details.
                try:
                    full_dataset = bq_service.datasets().get(projectId=project_id, datasetId=dataset_id).execute()
                    
                    # Compute our custom calculated fields
                    calculated_fields = calculate_bq_objects_and_owners(bq_service, project_id, dataset_id, full_dataset)
                    full_dataset['_calculated'] = calculated_fields
                    
                    datasets_data.append(full_dataset)
                except HttpError as e:
                    print(f"  -> Error fetching details for {dataset_id}: {e._get_reason()}")
                    
            request = bq_service.datasets().list_next(previous_request=request, previous_response=response)
        return datasets_data
    except HttpError as e:
        if e.resp.status in [403, 404]:
            pass  # Silently skip projects where BQ API is disabled or permissions are missing[cite: 6]
        else:
            print(f"API Error checking project {project_id}: {e._get_reason()}")
        return []

def write_to_excel(final_output, file_path):
    """Writes the parsed dictionary to an Excel file."""
    wb = Workbook()
    ws = wb.create_sheet(title="BigQuery Configs")
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    headers_written = False

    for instance_data in final_output:
        headers = []
        values = []

        for key, value in instance_data.items():
            headers.append(key)
            if isinstance(value, list):
                clean_value = "\n".join([str(i) for i in value])
                values.append(clean_value)
            elif isinstance(value, dict):
                clean_value = "\n".join([f"{k}: {v}" for k, v in value.items()])
                values.append(clean_value)
            else:
                values.append(value)

        if not headers_written:
            ws.append(headers)
            header_font = Font(bold=True)
            for cell in ws[1]:
                cell.font = header_font
            headers_written = True

        ws.append(values)

    wb.save(file_path)

def helper_function(data, path):
    keys = path.split('.')
    current = data
    for key in keys:
        if isinstance(current, dict):
            current = current.get(key, "N/A")
        else:
            return "N/A"
    return current

def get_nested_value(data, path):
    if isinstance(path, dict):
        return {label: helper_function(data, api_config) for label, api_config in path.items()}
    return helper_function(data, path)

def arrange_data_based_on_json_file(data, template_file):
    with open(template_file, 'r') as f:
        template = json.load(f)

    final_report = {}
    for label, api_config_path in template.items():
        final_report[label] = get_nested_value(data, api_config_path)

    return final_report

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="BigQuery Dataset Inventory script (All Accessible Projects)")
    parser.add_argument("--filename_prefix", type=str, required=True, help="Prefix for the output Excel filename")
    parser.add_argument("--mapping_file", type=str, required=False, default="bq_config.json", help="Path to the mapping JSON template")
    args = parser.parse_args()

    try:
        credentials, default_project = google.auth.default()
        crm_service = discovery.build('cloudresourcemanager', 'v3', credentials=credentials)
        bq_service = discovery.build('bigquery', 'v2', credentials=credentials)
    except Exception as e:
        print(f"Authentication failed: {e}")
        print("Run 'gcloud auth application-default login' first.")
        sys.exit(1)

    final_output = []
    timestamp = datetime.datetime.now().strftime("%Y%m%d%H%M%S")
    filename = f"{args.filename_prefix}_{timestamp}.xlsx"

    # 1. Fetch all projects your user/service account can see[cite: 6]
    projects = get_active_projects(crm_service)
    print(f"Found {len(projects)} active projects. Scanning for BigQuery datasets...")

    # 2. Iterate through projects and list datasets[cite: 6]
    for project_id in projects:
        datasets = get_datasets_for_project(bq_service, project_id)
        if datasets:
            print(f"Found {len(datasets)} BigQuery dataset(s) in project: {project_id}")
            for ds_data in datasets:
                arranged_data = arrange_data_based_on_json_file(ds_data, args.mapping_file)
                final_output.append(arranged_data)

    # 3. Write to Excel[cite: 6]
    if final_output:
        write_to_excel(final_output, filename)
        print(f"\nReport successfully written to: {filename}")
    else:
        print("\nNo BigQuery datasets found across any scanned projects. File creation skipped.")