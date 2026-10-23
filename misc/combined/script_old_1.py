import json
import argparse
import datetime
import sys
import subprocess
from openpyxl import Workbook
from openpyxl.styles import Font

from googleapiclient import discovery
from googleapiclient.errors import HttpError
import google.auth

# ==========================================
# RESOURCE MANAGER (PROJECT DISCOVERY)
# ==========================================
def get_org_projects(crm_service, org_id):
    print(f"\n[INFO] Discovering all active projects under Organization ID: {org_id}...")
    projects = []

    def walk_hierarchy(parent_name):
        try:
            request = crm_service.projects().list(parent=parent_name)
            while request is not None:
                response = request.execute()
                for project in response.get('projects', []):
                    if project.get('state') == 'ACTIVE':
                        projects.append(project.get('projectId'))
                request = crm_service.projects().list_next(previous_request=request, previous_response=response)
        except HttpError as e:
            print(f"Error fetching projects for {parent_name}: {e}")

        try:
            request = crm_service.folders().list(parent=parent_name)
            while request is not None:
                response = request.execute()
                for folder in response.get('folders', []):
                    if folder.get('state') == 'ACTIVE':
                        walk_hierarchy(folder.get('name'))
                request = crm_service.folders().list_next(previous_request=request, previous_response=response)
        except HttpError as e:
            print(f"Error fetching folders for {parent_name}: {e}")

    walk_hierarchy(f"organizations/{org_id}")
    return projects


# ==========================================
# SERVICE FETCHING FUNCTIONS
# ==========================================

# 1. Spanner
def list_spanner_instances(spanner_service, project_id):
    instances, parent = [], f"projects/{project_id}"
    try:
        request = spanner_service.projects().instances().list(parent=parent)
        while request is not None:
            response = request.execute()
            if 'instances' in response: instances.extend(response['instances'])
            request = spanner_service.projects().instances().list_next(previous_request=request, previous_response=response)
    except HttpError: pass
    return instances

# 2. Redis
def list_redis_instances(redis_service, project_id):
    instances, parent = [], f"projects/{project_id}/locations/-"
    try:
        request = redis_service.projects().locations().instances().list(parent=parent)
        while request is not None:
            response = request.execute()
            if 'instances' in response: instances.extend(response['instances'])
            request = redis_service.projects().locations().instances().list_next(previous_request=request, previous_response=response)
    except HttpError: pass
    return instances

# 3. Cloud SQL
def list_cloudsql_instances(sql_service, project_id):
    instances = []
    try:
        request = sql_service.instances().list(project=project_id)
        while request is not None:
            response = request.execute()
            if 'items' in response: instances.extend(response['items'])
            request = sql_service.instances().list_next(previous_request=request, previous_response=response)
    except HttpError: pass
    return instances

def enrich_sql_calculations(data):
    calc = {}
    tier = data.get("settings", {}).get("tier", "")
    if tier.startswith("db-custom-"):
        parts = tier.split("-")
        if len(parts) >= 4:
            calc["vcpu"] = parts[2]
            calc["memory_gb"] = str(int(parts[3]) / 1024)
    else:
        calc["vcpu"] = "Tied to Tier"
        calc["memory_gb"] = "Tied to Tier"
    calc["replica_count"] = len(data.get("replicaNames", []))
    calc["replica_names"] = data.get("replicaNames", [])
    data["_calculated"] = calc
    return data

# 4. BigQuery
def calculate_bq_objects_and_owners(bq_service, project_id, dataset_id, dataset_details):
    owners = []
    for entry in dataset_details.get("access", []):
        if entry.get("role") == "OWNER":
            owner = entry.get("userByEmail") or entry.get("specialGroup") or entry.get("groupByEmail")
            if owner: owners.append(owner)
    
    table_count, view_count, routine_count, model_count = 0, 0, 0, 0
    try:
        req = bq_service.tables().list(projectId=project_id, datasetId=dataset_id)
        while req is not None:
            resp = req.execute()
            for obj in resp.get('tables', []):
                t = obj.get("type", "")
                if t in ["TABLE", "EXTERNAL", "SNAPSHOT", "CLONE"]: table_count += 1
                elif t in ["VIEW", "MATERIALIZED_VIEW"]: view_count += 1
            req = bq_service.tables().list_next(previous_request=req, previous_response=resp)
    except HttpError: pass 
    return {
        "dataset_owners": ", ".join(owners) if owners else "None",
        "table_count": table_count,
        "view_count": view_count,
        "routine_count": 0, 
        "model_count": 0,   
        "total_object_count": table_count + view_count
    }

def get_datasets_for_project(bq_service, project_id):
    datasets_data = []
    try:
        request = bq_service.datasets().list(projectId=project_id)
        while request is not None:
            response = request.execute()
            for ds in response.get('datasets', []):
                dataset_id = ds['datasetReference']['datasetId']
                try:
                    full_dataset = bq_service.datasets().get(projectId=project_id, datasetId=dataset_id).execute()
                    full_dataset['_calculated'] = calculate_bq_objects_and_owners(bq_service, project_id, dataset_id, full_dataset)
                    datasets_data.append(full_dataset)
                except HttpError: pass
            request = bq_service.datasets().list_next(previous_request=request, previous_response=response)
    except HttpError: pass
    return datasets_data

# 5. AlloyDB
def list_alloydb_clusters(alloydb_service, project_id):
    clusters, parent = [], f"projects/{project_id}/locations/-"
    try:
        request = alloydb_service.projects().locations().clusters().list(parent=parent)
        while request is not None:
            response = request.execute()
            if 'clusters' in response: clusters.extend(response['clusters'])
            request = alloydb_service.projects().locations().clusters().list_next(previous_request=request, previous_response=response)
    except HttpError: pass
    return clusters

def list_alloydb_instances(alloydb_service, cluster_name):
    instances = []
    try:
        request = alloydb_service.projects().locations().clusters().instances().list(parent=cluster_name)
        while request is not None:
            response = request.execute()
            if 'instances' in response: instances.extend(response['instances'])
            request = alloydb_service.projects().locations().clusters().instances().list_next(previous_request=request, previous_response=response)
    except HttpError: pass
    return instances

# 6. Memcached
def list_memcached_instances(memcache_service, project_id):
    instances, parent = [], f"projects/{project_id}/locations/-"
    try:
        request = memcache_service.projects().locations().instances().list(parent=parent)
        while request is not None:
            response = request.execute()
            if 'instances' in response: instances.extend(response['instances'])
            request = memcache_service.projects().locations().instances().list_next(previous_request=request, previous_response=response)
    except HttpError: pass
    return instances

# 7. Valkey (Via Subprocess/REST API due to client library limitation)
def list_valkey_instances(project_id):
    try:
        token = subprocess.check_output(["gcloud", "auth", "print-access-token"], text=True).strip()
        url = f"https://memorystore.googleapis.com/v1/projects/{project_id}/locations/-/instances"
        curl_cmd = ["curl", "-s", "-H", f"Authorization: Bearer {token}", url]
        
        result = subprocess.check_output(curl_cmd, text=True)
        response = json.loads(result)
        
        if 'error' in response:
            return []
            
        all_instances = response.get('instances', [])
        return [inst for inst in all_instances if "VALKEY" in inst.get("engineVersion", "")]
    except Exception as e:
        return []


# ==========================================
# UTILITIES & DATA ARRANGEMENT
# ==========================================
def helper_function(data, path):
    keys = path.split('.')
    current = data
    for idx, key in enumerate(keys):
        if isinstance(current, dict):
            current = current.get(key, "N/A")
        elif isinstance(current, list):
            remaining_path = '.'.join(keys[idx:])
            extracted_list = [helper_function(item, remaining_path) for item in current if helper_function(item, remaining_path) != "N/A"]
            return extracted_list if extracted_list else "N/A"
        else: return "N/A"
        
    if isinstance(current, str) and "projects/" in current and "locations/" in current:
        if path == "name": return current.split('/')[-1]
    return current

def get_nested_value(data, path):
    if isinstance(path, dict):
        return {label: helper_function(data, api_config) for label, api_config in path.items()}
    return helper_function(data, path)

def arrange_data_based_on_json(data, template):
    final_report = {}
    if "name" in data:
        raw_name = data.get("name", "")
        data["_calculated_short_name"] = raw_name.split('/')[-1] if '/' in raw_name else raw_name
    if "kind" in data and data["kind"] == "sql#instance":
        data = enrich_sql_calculations(data)

    for label, api_config_path in template.items():
        if api_config_path == "project_id": continue
        if api_config_path == "name" and "_calculated_short_name" in data:
            final_report[label] = data.get("_calculated_short_name", "N/A")
        else:
            final_report[label] = get_nested_value(data, api_config_path)
    return final_report


# ==========================================
# EXCEL EXPORT
# ==========================================
def write_to_excel(service_outputs, file_path):
    wb = Workbook()
    if "Sheet" in wb.sheetnames: del wb["Sheet"]

    for sheet_name, data in service_outputs.items():
        if not data: continue
        ws = wb.create_sheet(title=sheet_name)
        headers_written = False
        for instance_data in data:
            headers, values = [], []
            for key, value in instance_data.items():
                headers.append(key)
                if isinstance(value, list): values.append("\n".join([str(i) for i in value]))
                elif isinstance(value, dict): values.append("\n".join([f"{k}: {v}" for k, v in value.items()]))
                else: values.append(value)
                
            if not headers_written:
                ws.append(headers)
                for cell in ws[1]: cell.font = Font(bold=True)
                headers_written = True
            ws.append(values)
            
    if len(wb.sheetnames) == 0: wb.create_sheet("No Data")
    wb.save(file_path)


# ==========================================
# MAIN EXECUTION
# ==========================================
if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Unified GCP DB Inventory Script")
    parser.add_argument("--org_id", type=str, required=True, help="GCP Organization ID to scan")
    parser.add_argument("--filename_prefix", type=str, required=True, help="Prefix for Excel filename")
    parser.add_argument("--mapping_file", type=str, required=False, default="combined_config.json")
    args = parser.parse_args()

    try:
        credentials, _ = google.auth.default()
        crm_service = discovery.build('cloudresourcemanager', 'v3', credentials=credentials)
        redis_service = discovery.build('redis', 'v1', credentials=credentials)
        sql_service = discovery.build('sqladmin', 'v1beta4', credentials=credentials)
        bq_service = discovery.build('bigquery', 'v2', credentials=credentials)
        spanner_service = discovery.build('spanner', 'v1', credentials=credentials)
        alloydb_service = discovery.build('alloydb', 'v1', credentials=credentials)
        memcache_service = discovery.build('memcache', 'v1', credentials=credentials)
        # removed memorystore_service
    except Exception as e:
        print(f"Authentication failed: {e}")
        sys.exit(1)

    with open(args.mapping_file, 'r') as f:
        master_template = json.load(f)

    outputs = {
        "redis": [], "cloudsql": [], "bigquery": [], 
        "spanner": [], "alloydb": [], "memcached": [], "valkey": []
    }
    
    projects = get_org_projects(crm_service, args.org_id)

    for i, project_id in enumerate(projects, 1):
        print(f"[{i}/{len(projects)}] Scanning Project: {project_id}")
        
        # 1. Redis
        for inst in list_redis_instances(redis_service, project_id):
            row = {"Project ID": project_id}
            row.update(arrange_data_based_on_json(inst, master_template.get("redis", {})))
            outputs["redis"].append(row)

        # 2. Cloud SQL
        for inst in list_cloudsql_instances(sql_service, project_id):
            row = {"Project ID": project_id}
            row.update(arrange_data_based_on_json(inst, master_template.get("cloudsql", {})))
            outputs["cloudsql"].append(row)

        # 3. BigQuery
        for ds in get_datasets_for_project(bq_service, project_id):
            row = {"Project ID": project_id}
            row.update(arrange_data_based_on_json(ds, master_template.get("bigquery", {})))
            outputs["bigquery"].append(row)

        # 4. Spanner
        for inst in list_spanner_instances(spanner_service, project_id):
            row = {"Project ID": project_id}
            row.update(arrange_data_based_on_json(inst, master_template.get("spanner", {})))
            outputs["spanner"].append(row)
            
        # 5. AlloyDB
        for cluster in list_alloydb_clusters(alloydb_service, project_id):
            cluster_name = cluster.get("name", "")
            for inst in list_alloydb_instances(alloydb_service, cluster_name):
                combined = {"project_id": project_id, "cluster": cluster, "instance": inst}
                row = {"Project ID": project_id}
                row.update(arrange_data_based_on_json(combined, master_template.get("alloydb", {}).get("cluster_mappings", {})))
                row.update(arrange_data_based_on_json(combined, master_template.get("alloydb", {}).get("instance_mappings", {})))
                outputs["alloydb"].append(row)

        # 6. Memcached
        for inst in list_memcached_instances(memcache_service, project_id):
            row = {"Project ID": project_id}
            row.update(arrange_data_based_on_json(inst, master_template.get("memcached", {})))
            outputs["memcached"].append(row)

        # 7. Valkey (Subprocess)
        for inst in list_valkey_instances(project_id):
            row = {"Project ID": project_id}
            row.update(arrange_data_based_on_json(inst, master_template.get("valkey", {})))
            outputs["valkey"].append(row)

    filename = f"{args.filename_prefix}_{datetime.datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    write_to_excel(outputs, filename)
    print(f"\n[DONE] Report written to: {filename}")