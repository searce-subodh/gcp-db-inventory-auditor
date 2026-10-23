import json
import argparse
import datetime
import sys
from openpyxl import Workbook
from openpyxl.styles import Font

from googleapiclient import discovery
from googleapiclient.errors import HttpError
import google.auth

# ==========================================
# RESOURCE MANAGER (PROJECT DISCOVERY)
# ==========================================
def get_org_projects(crm_service, org_id):
    """Recursively fetches all active projects under an Organization (including inside folders)[cite: 7]."""
    print(f"\n[INFO] Discovering all active projects under Organization ID: {org_id}...")
    projects = []

    def walk_hierarchy(parent_name):
        try:
            request = crm_service.projects().list(parent=parent_name)
            while request is not None:
                response = request.execute()
                for project in response.get('projects', []):
                    if project.get('state') == 'ACTIVE':
                        projects.append(project.get('projectId'))
                request = crm_service.projects().list_next(previous_request=request, previous_response=response)
        except HttpError as e:
            print(f"Error fetching projects for {parent_name}: {e}")

        try:
            request = crm_service.folders().list(parent=parent_name)
            while request is not None:
                response = request.execute()
                for folder in response.get('folders', []):
                    if folder.get('state') == 'ACTIVE':
                        walk_hierarchy(folder.get('name'))
                request = crm_service.folders().list_next(previous_request=request, previous_response=response)
        except HttpError as e:
            print(f"Error fetching folders for {parent_name}: {e}")

    walk_hierarchy(f"organizations/{org_id}")
    return projects

# ==========================================
# SPANNER INVENTORY
# ==========================================
def list_spanner_instances(spanner_service, project_id):
    """Lists all Spanner instances in a project[cite: 6]."""
    instances = []
    parent = f"projects/{project_id}"
    
    try:
        # Note: Spanner instances are parented directly to the project, not locations[cite: 6].
        request = spanner_service.projects().instances().list(parent=parent)
        while request is not None:
            response = request.execute()
            if 'instances' in response:
                instances.extend(response['instances'])
            request = spanner_service.projects().instances().list_next(previous_request=request, previous_response=response)
        return instances
    except HttpError as e:
        if e.resp.status not in [403, 404]:
            print(f"  [WARN] Spanner API Error checking project {project_id}: {e._get_reason()}")
        return []

# ==========================================
# REDIS INVENTORY
# ==========================================
def list_redis_instances(redis_service, project_id):
    instances = []
    parent = f"projects/{project_id}/locations/-"
    try:
        request = redis_service.projects().locations().instances().list(parent=parent)
        while request is not None:
            response = request.execute()
            if 'instances' in response:
                instances.extend(response['instances'])
            request = redis_service.projects().locations().instances().list_next(previous_request=request, previous_response=response)
        return instances
    except HttpError as e:
        if e.resp.status not in [403, 404]:
            print(f"  [WARN] Redis API Error checking project {project_id}: {e._get_reason()}")
        return []

# ==========================================
# CLOUD SQL INVENTORY
# ==========================================
def list_cloudsql_instances(sql_service, project_id):
    instances = []
    try:
        request = sql_service.instances().list(project=project_id)
        while request is not None:
            response = request.execute()
            if 'items' in response:
                instances.extend(response['items'])
            request = sql_service.instances().list_next(previous_request=request, previous_response=response)
        return instances
    except HttpError as e:
        if e.resp.status not in [403, 404]:
            print(f"  [WARN] SQL API Error checking project {project_id}: {e._get_reason()}")
        return []

def enrich_sql_calculations(data):
    """Calculates custom fields specific to Cloud SQL[cite: 7]."""
    calc = {}
    tier = data.get("settings", {}).get("tier", "")
    if tier.startswith("db-custom-"):
        parts = tier.split("-")
        if len(parts) >= 4:
            calc["vcpu"] = parts[2]
            calc["memory_gb"] = str(int(parts[3]) / 1024)
    else:
        calc["vcpu"] = "Tied to Tier"
        calc["memory_gb"] = "Tied to Tier"

    replica_names = data.get("replicaNames", [])
    calc["replica_count"] = len(replica_names)
    calc["replica_names"] = replica_names

    data["_calculated"] = calc
    return data

# ==========================================
# BIGQUERY INVENTORY
# ==========================================
def calculate_bq_objects_and_owners(bq_service, project_id, dataset_id, dataset_details):
    """Calculates object counts and parses owners to populate the _calculated fields[cite: 7]."""
    owners = []
    access_list = dataset_details.get("access", [])
    for entry in access_list:
        if entry.get("role") == "OWNER":
            owner = entry.get("userByEmail") or entry.get("specialGroup") or entry.get("groupByEmail")
            if owner:
                owners.append(owner)
    
    dataset_owners = ", ".join(owners) if owners else "None"
    table_count, view_count, routine_count, model_count = 0, 0, 0, 0

    try:
        tables_req = bq_service.tables().list(projectId=project_id, datasetId=dataset_id)
        while tables_req is not None:
            tables_resp = tables_req.execute()
            for obj in tables_resp.get('tables', []):
                obj_type = obj.get("type", "")
                if obj_type in ["TABLE", "EXTERNAL", "SNAPSHOT", "CLONE"]:
                    table_count += 1
                elif obj_type in ["VIEW", "MATERIALIZED_VIEW"]:
                    view_count += 1
            tables_req = bq_service.tables().list_next(previous_request=tables_req, previous_response=tables_resp)

        routines_req = bq_service.routines().list(projectId=project_id, datasetId=dataset_id)
        while routines_req is not None:
            routines_resp = routines_req.execute()
            routine_count += len(routines_resp.get('routines', []))
            routines_req = bq_service.routines().list_next(previous_request=routines_req, previous_response=routines_resp)

        models_req = bq_service.models().list(projectId=project_id, datasetId=dataset_id)
        while models_req is not None:
            models_resp = models_req.execute()
            model_count += len(models_resp.get('models', []))
            models_req = bq_service.models().list_next(previous_request=models_req, previous_response=models_resp)

    except HttpError as e:
        pass 

    return {
        "dataset_owners": dataset_owners,
        "table_count": table_count,
        "view_count": view_count,
        "routine_count": routine_count,
        "model_count": model_count,
        "total_object_count": table_count + view_count + routine_count + model_count
    }

def get_datasets_for_project(bq_service, project_id):
    """Lists all datasets in a project and fetches their full metadata[cite: 7]."""
    datasets_data = []
    try:
        request = bq_service.datasets().list(projectId=project_id)
        while request is not None:
            response = request.execute()
            for ds in response.get('datasets', []):
                dataset_id = ds['datasetReference']['datasetId']
                try:
                    full_dataset = bq_service.datasets().get(projectId=project_id, datasetId=dataset_id).execute()
                    calculated_fields = calculate_bq_objects_and_owners(bq_service, project_id, dataset_id, full_dataset)
                    full_dataset['_calculated'] = calculated_fields
                    datasets_data.append(full_dataset)
                except HttpError as e:
                    print(f"  -> Error fetching BQ details for {dataset_id}: {e._get_reason()}")
                    
            request = bq_service.datasets().list_next(previous_request=request, previous_response=response)
        return datasets_data
    except HttpError as e:
        if e.resp.status not in [403, 404]:
            print(f"  [WARN] BQ API Error checking project {project_id}: {e._get_reason()}")
        return []

# ==========================================
# UTILITIES & EXCEL EXPORT
# ==========================================
def write_to_excel(redis_output, sql_output, bq_output, spanner_output, file_path):
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    def write_sheet(data, sheet_name):
        if not data: return
        ws = wb.create_sheet(title=sheet_name)
        headers_written = False
        for instance_data in data:
            headers = []
            values = []
            for key, value in instance_data.items():
                headers.append(key)
                if isinstance(value, list):
                    values.append("\n".join([str(i) for i in value]))
                elif isinstance(value, dict):
                    values.append("\n".join([f"{k}: {v}" for k, v in value.items()]))
                else:
                    values.append(value)
            if not headers_written:
                ws.append(headers)
                for cell in ws[1]: cell.font = Font(bold=True)
                headers_written = True
            ws.append(values)

    write_sheet(redis_output, "Redis Configs")
    write_sheet(sql_output, "Cloud SQL Configs")
    write_sheet(bq_output, "BigQuery Datasets")
    write_sheet(spanner_output, "Spanner Configs")
    wb.save(file_path)

def helper_function(data, path):
    keys = path.split('.')
    current = data
    for key in keys:
        if isinstance(current, dict):
            current = current.get(key, "N/A")
        else:
            return "N/A"
    return current

def get_nested_value(data, path):
    if isinstance(path, dict):
        return {label: helper_function(data, api_config) for label, api_config in path.items()}
    return helper_function(data, path)

def arrange_data_based_on_json(data, template):
    final_report = {}
    
    if "name" in data:
        raw_name = data.get("name", "")
        data["_calculated_short_name"] = raw_name.split('/')[-1] if '/' in raw_name else raw_name

    if "kind" in data and data["kind"] == "sql#instance":
        data = enrich_sql_calculations(data)

    for label, api_config_path in template.items():
        if api_config_path == "name" and "_calculated_short_name" in data:
            final_report[label] = data.get("_calculated_short_name", "N/A")
        else:
            final_report[label] = get_nested_value(data, api_config_path)

    return final_report

# ==========================================
# MAIN EXECUTION
# ==========================================
if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="GCP Inventory (Redis, Cloud SQL, BigQuery, Spanner) by Organization")
    parser.add_argument("--org_id", type=str, required=True, help="GCP Organization ID (numeric string) to scan")
    parser.add_argument("--filename_prefix", type=str, required=True, help="Prefix for Excel filename")
    parser.add_argument("--mapping_file", type=str, required=False, default="combined_config.json", help="Path to unified JSON template")
    args = parser.parse_args()

    try:
        credentials, _ = google.auth.default()
        crm_service = discovery.build('cloudresourcemanager', 'v3', credentials=credentials)
        redis_service = discovery.build('redis', 'v1', credentials=credentials)
        sql_service = discovery.build('sqladmin', 'v1beta4', credentials=credentials)
        bq_service = discovery.build('bigquery', 'v2', credentials=credentials)
        spanner_service = discovery.build('spanner', 'v1', credentials=credentials)
    except Exception as e:
        print(f"Authentication failed: {e}")
        sys.exit(1)

    with open(args.mapping_file, 'r') as f:
        master_template = json.load(f)

    redis_template = master_template.get("redis", {})
    sql_template = master_template.get("cloudsql", {})
    bq_template = master_template.get("bigquery", {})
    spanner_template = master_template.get("spanner", {})

    redis_output, sql_output, bq_output, spanner_output = [], [], [], []
    
    projects = get_org_projects(crm_service, args.org_id)
    print(f"[INFO] Found {len(projects)} active projects under Organization {args.org_id}. Starting scan...\n")

    for i, project_id in enumerate(projects, 1):
        print(f"[{i}/{len(projects)}] Scanning Project: {project_id}")
        
        # 1. Redis Scan
        print(f"  -> Fetching Redis instances...")
        r_instances = list_redis_instances(redis_service, project_id)
        if r_instances:
            print(f"  -> SUCCESS: Found {len(r_instances)} Redis instance(s).")
            for data in r_instances:
                arranged = arrange_data_based_on_json(data, redis_template)
                arranged["Project ID"] = project_id
                redis_output.append(arranged)

        # 2. Cloud SQL Scan
        print(f"  -> Fetching Cloud SQL instances...")
        s_instances = list_cloudsql_instances(sql_service, project_id)
        if s_instances:
            print(f"  -> SUCCESS: Found {len(s_instances)} Cloud SQL instance(s).")
            for data in s_instances:
                arranged = arrange_data_based_on_json(data, sql_template)
                arranged["Project ID"] = project_id 
                sql_output.append(arranged)
                
        # 3. BigQuery Scan
        print(f"  -> Fetching BigQuery datasets...")
        b_datasets = get_datasets_for_project(bq_service, project_id)
        if b_datasets:
            print(f"  -> SUCCESS: Found {len(b_datasets)} BigQuery dataset(s).")
            for data in b_datasets:
                arranged = arrange_data_based_on_json(data, bq_template)
                arranged["Project ID"] = project_id 
                bq_output.append(arranged)

        # 4. Spanner Scan
        print(f"  -> Fetching Spanner instances...")
        sp_instances = list_spanner_instances(spanner_service, project_id)
        if sp_instances:
            print(f"  -> SUCCESS: Found {len(sp_instances)} Spanner instance(s).")
            for data in sp_instances:
                arranged = arrange_data_based_on_json(data, spanner_template)
                arranged["Project ID"] = project_id
                spanner_output.append(arranged)
        
        print("-" * 40)

    filename = f"{args.filename_prefix}_{datetime.datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    write_to_excel(redis_output, sql_output, bq_output, spanner_output, filename)
    print(f"\n[DONE] Report written to: {filename}")
    print(f"Total extracted: {len(redis_output)} Redis, {len(sql_output)} Cloud SQL, {len(bq_output)} BigQuery Datasets, {len(spanner_output)} Spanner Instances.")