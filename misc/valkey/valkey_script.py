import json
import argparse
import datetime
import sys
from openpyxl import Workbook
from openpyxl.styles import Font

from googleapiclient import discovery
from googleapiclient.errors import HttpError
import google.auth

def get_active_projects(crm_service):
    """Fetches all active projects accessible by the authenticated credentials."""
    print("Discovering all active projects...")
    projects = []
    query = "state:ACTIVE"
    
    try:
        request = crm_service.projects().search(query=query)
        while request is not None:
            response = request.execute()
            for project in response.get('projects', []):
                projects.append(project.get('projectId'))
            request = crm_service.projects().search_next(previous_request=request, previous_response=response)
        return projects
    except HttpError as e:
        print(f"Error fetching projects: {e}")
        sys.exit(1)

def list_valkey_instances(memorystore_service, project_id):
    """Lists all Valkey/Memorystore instances in a project across all regions."""
    instances = []
    parent = f"projects/{project_id}/locations/-"
    
    try:
        request = memorystore_service.projects().locations().instances().list(parent=parent)
        while request is not None:
            response = request.execute()
            if 'instances' in response:
                instances.extend(response['instances'])
            request = memorystore_service.projects().locations().instances().list_next(previous_request=request, previous_response=response)
        return instances
    except HttpError as e:
        # Silently skip projects where the Memorystore API is disabled or permissions are missing
        if e.resp.status in [403, 404]:
            pass 
        else:
            print(f"API Error checking project {project_id}: {e._get_reason()}")
        return []

def write_to_excel(final_output, file_path):
    """Writes the parsed dictionary to an Excel file."""
    wb = Workbook()
    ws = wb.create_sheet(title="Valkey Inventory")
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    headers_written = False

    for instance_data in final_output:
        headers = []
        values = []

        for key, value in instance_data.items():
            headers.append(key)
            if isinstance(value, list):
                clean_value = "\n".join([str(i) for i in value])
                values.append(clean_value)
            elif isinstance(value, dict):
                clean_value = "\n".join([f"{k}: {v}" for k, v in value.items()])
                values.append(clean_value)
            else:
                values.append(value)

        if not headers_written:
            ws.append(headers)
            header_font = Font(bold=True)
            for cell in ws[1]:
                cell.font = header_font
            headers_written = True

        ws.append(values)

    wb.save(file_path)

def helper_function(data, path):
    """Traverses the dictionary based on the dot-notation path, handling lists elegantly."""
    keys = path.split('.')
    current = data
    
    for idx, key in enumerate(keys):
        if isinstance(current, dict):
            current = current.get(key, "N/A")
        elif isinstance(current, list):
            remaining_path = '.'.join(keys[idx:])
            extracted_list = []
            for item in current:
                val = helper_function(item, remaining_path)
                if val != "N/A":
                    extracted_list.append(val)
            return extracted_list if extracted_list else "N/A"
        else:
            return "N/A"
    
    # Extract just the short name for GCP resource paths
    if isinstance(current, str) and "projects/" in current and "locations/" in current:
        if path == "name":
            return current.split('/')[-1]
            
    return current

def get_nested_value(data, path):
    if isinstance(path, dict):
        return {label: helper_function(data, api_config) for label, api_config in path.items()}
    return helper_function(data, path)

def arrange_data_based_on_json_file(instance_data, template_file):
    """Arranges data based on the JSON mapping file."""
    with open(template_file, 'r') as f:
        template = json.load(f)

    final_report = {}
    
    for label, api_config_path in template.items():
        # Skip Project ID here, it gets injected in the main loop
        if api_config_path == "project_id":
            continue
        final_report[label] = get_nested_value(instance_data, api_config_path)

    return final_report

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Valkey/Memorystore Inventory Fetching Script")
    parser.add_argument("--filename_prefix", type=str, required=True, help="Prefix for the output Excel filename")
    parser.add_argument("--mapping_file", type=str, required=False, default="valkey_config.json", help="Path to the mapping JSON template")
    args = parser.parse_args()



    try:
        credentials, default_project = google.auth.default()
        crm_service = discovery.build('cloudresourcemanager', 'v3', credentials=credentials)
        memorystore_service = discovery.build('memorystore', 'v1', credentials=credentials)
    except Exception as e:
        print(f"Failed to initialize Google API clients: {e}")
        sys.exit(1)

    final_output = []
    timestamp = datetime.datetime.now().strftime("%Y%m%d%H%M%S")
    filename = f"{args.filename_prefix}_{timestamp}.xlsx"

    projects = get_active_projects(crm_service)
    print(f"Found {len(projects)} active projects. Scanning for Valkey instances...")

    for project_id in projects:
        instances = list_valkey_instances(memorystore_service, project_id)
        if instances:
            # Filter specifically for Valkey engines if needed, though the API returns all Memorystore instances
            valkey_instances = [inst for inst in instances if "VALKEY" in inst.get("engineVersion", "")]
            
            if valkey_instances:
                print(f"Found {len(valkey_instances)} Valkey instance(s) in project: {project_id}")
                
                for instance in valkey_instances:
                    arranged_data = arrange_data_based_on_json_file(instance, args.mapping_file)
                    
                    # Force inject Project ID at the start of the row
                    final_row = {"Project ID": project_id}
                    final_row.update(arranged_data)
                    
                    final_output.append(final_row)

    if final_output:
        write_to_excel(final_output, filename)
        print(f"\nReport successfully written to: {filename}")
    else:
        print("\nNo Valkey instances found across any scanned projects. File creation skipped.")