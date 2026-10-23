import json
import argparse
import datetime
import sys
from openpyxl import Workbook
from openpyxl.styles import Font

from googleapiclient import discovery
from googleapiclient.errors import HttpError
import google.auth

def get_active_projects(crm_service):
    """Fetches all active projects accessible by the authenticated credentials."""
    print("Discovering all active projects you have access to...")
    projects = []
    query = "state:ACTIVE"
    
    try:
        request = crm_service.projects().search(query=query)
        while request is not None:
            response = request.execute()
            for project in response.get('projects', []):
                projects.append(project.get('projectId'))
            request = crm_service.projects().search_next(previous_request=request, previous_response=response)
        return projects
    except HttpError as e:
        print(f"Error fetching projects: {e}")
        sys.exit(1)

def list_memcached_instances(memcache_service, project_id):
    """Lists all Memcached instances in a project across all regions using the '-' wildcard."""
    instances = []
    parent = f"projects/{project_id}/locations/-"
    
    try:
        request = memcache_service.projects().locations().instances().list(parent=parent)
        while request is not None:
            response = request.execute()
            if 'instances' in response:
                instances.extend(response['instances'])
            request = memcache_service.projects().locations().instances().list_next(previous_request=request, previous_response=response)
        return instances
    except HttpError as e:
        # Silently skip projects where the Memcached API is disabled or permissions are missing
        if e.resp.status in [403, 404]:
            pass 
        else:
            print(f"API Error checking project {project_id}: {e._get_reason()}")
        return []

def write_to_excel(final_output, file_path):
    """Writes the parsed dictionary to an Excel file."""
    wb = Workbook()
    ws = wb.create_sheet(title="Memcached Configs")
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    headers_written = False

    for instance_data in final_output:
        headers = []
        values = []

        for key, value in instance_data.items():
            headers.append(key)
            if isinstance(value, list):
                clean_value = "\n".join([str(i) for i in value])
                values.append(clean_value)
            elif isinstance(value, dict):
                clean_value = "\n".join([f"{k}: {v}" for k, v in value.items()])
                values.append(clean_value)
            else:
                values.append(value)

        if not headers_written:
            ws.append(headers)
            header_font = Font(bold=True)
            for cell in ws[1]:
                cell.font = header_font
            headers_written = True

        ws.append(values)

    wb.save(file_path)

def helper_function(data, path):
    """Traverses the dictionary based on dot-notation, handling lists elegantly."""
    keys = path.split('.')
    current = data
    
    for idx, key in enumerate(keys):
        if isinstance(current, dict):
            current = current.get(key, "N/A")
        elif isinstance(current, list):
            remaining_path = '.'.join(keys[idx:])
            extracted_list = []
            for item in current:
                val = helper_function(item, remaining_path)
                if val != "N/A":
                    extracted_list.append(val)
            return extracted_list if extracted_list else "N/A"
        else:
            return "N/A"
    return current

def get_nested_value(data, path):
    if isinstance(path, dict):
        return {label: helper_function(data, api_config) for label, api_config in path.items()}
    return helper_function(data, path)

def arrange_data_based_on_json_file(data, template_file):
    with open(template_file, 'r') as f:
        template = json.load(f)

    final_report = {}
    
    # Extract short name from the long GCP resource path
    raw_name = data.get("name", "")
    short_name = raw_name.split('/')[-1] if '/' in raw_name else raw_name
    data["_calculated_short_name"] = short_name

    for label, api_config_path in template.items():
        if api_config_path == "name":
            final_report[label] = data.get("_calculated_short_name", "N/A")
        else:
            final_report[label] = get_nested_value(data, api_config_path)

    return final_report

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Memcached Instance Configs Listing Script")
    parser.add_argument("--filename_prefix", type=str, required=True, help="Prefix for the output Excel filename")
    parser.add_argument("--mapping_file", type=str, required=False, default="memcached_config.json", help="Path to the mapping JSON template")
    args = parser.parse_args()

    try:
        credentials, default_project = google.auth.default()
        crm_service = discovery.build('cloudresourcemanager', 'v3', credentials=credentials)
        memcache_service = discovery.build('memcache', 'v1', credentials=credentials)
    except Exception as e:
        print(f"Authentication failed: {e}")
        print("Run 'gcloud auth application-default login' first.")
        sys.exit(1)

    final_output = []
    timestamp = datetime.datetime.now().strftime("%Y%m%d%H%M%S")
    filename = f"{args.filename_prefix}_{timestamp}.xlsx"

    projects = get_active_projects(crm_service)
    print(f"Found {len(projects)} active projects. Scanning for Memcached instances...")

    for project_id in projects:
        instances = list_memcached_instances(memcache_service, project_id)
        if instances:
            print(f"Found {len(instances)} Memcached instance(s) in project: {project_id}")
            for instance_data in instances:
                arranged_data = arrange_data_based_on_json_file(instance_data, args.mapping_file)
                
                # Prepend Project ID to the row output
                final_row = {"Project ID": project_id}
                final_row.update(arranged_data)
                
                final_output.append(final_row)

    if final_output:
        write_to_excel(final_output, filename)
        print(f"\nReport successfully written to: {filename}")
    else:
        print("\nNo Memcached instances found across any scanned projects. File creation skipped.")